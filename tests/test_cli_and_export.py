import csv

import httpx
from openpyxl import load_workbook

from scraper.cli import main
from scraper.client import ClientConfig, Fetcher
from scraper.core import crawl
from scraper.exporters import export
from scraper.parsers import parse_quotes


def quotes_fetcher(fixture_html):
    pages = {
        "/": fixture_html("quotes_page1.html"),
        "/page/2/": fixture_html("quotes_page2.html"),
    }

    def handler(request):
        return httpx.Response(200, text=pages[request.url.path])

    return Fetcher(ClientConfig(delay=0), transport=httpx.MockTransport(handler), sleep=lambda s: None)


def test_crawl_follows_pagination(fixture_html):
    rows = list(crawl(quotes_fetcher(fixture_html), parse_quotes, "https://quotes.toscrape.com/"))
    assert [r["author"] for r in rows] == ["Albert Einstein", "J.K. Rowling", "Steve Martin"]


def test_crawl_limits(fixture_html):
    f = quotes_fetcher(fixture_html)
    assert len(list(crawl(f, parse_quotes, "https://quotes.toscrape.com/", max_pages=1))) == 2
    assert len(list(crawl(f, parse_quotes, "https://quotes.toscrape.com/", max_items=1))) == 1


def test_export_csv_and_xlsx(tmp_path):
    rows = [{"a": 1, "b": "中文"}, {"a": 2, "c": "x"}]
    p = export(rows, tmp_path / "out.csv")
    with p.open(encoding="utf-8-sig") as f:
        data = list(csv.DictReader(f))
    assert data[0]["b"] == "中文" and list(data[0]) == ["a", "b", "c"]

    x = export(rows, tmp_path / "out.xlsx")
    ws = load_workbook(x).active
    assert [c.value for c in ws[1]] == ["a", "b", "c"]
    assert ws["B2"].value == "中文"


def test_export_rejects_unknown(tmp_path):
    import pytest

    with pytest.raises(ValueError):
        export([], tmp_path / "out.txt")


def test_cli_end_to_end(fixture_html, tmp_path, capsys):
    out = tmp_path / "quotes.xlsx"
    rc = main(["quotes", "-o", str(out)], fetcher=quotes_fetcher(fixture_html))
    assert rc == 0
    assert "saved 3 rows" in capsys.readouterr().out
    assert load_workbook(out).active.max_row == 4
