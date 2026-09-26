"""Export scraped rows to CSV / Excel / JSON."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def _columns(rows: list[dict]) -> list[str]:
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


def to_csv(rows: Iterable[dict], path: str | Path) -> Path:
    rows = list(rows)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so Excel on Windows opens Chinese text correctly
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=_columns(rows))
        writer.writeheader()
        writer.writerows(rows)
    return path


def to_excel(rows: Iterable[dict], path: str | Path, sheet: str = "data") -> Path:
    rows = list(rows)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = _columns(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, start=1):
        width = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
    wb.save(path)
    return path


def to_json(rows: Iterable[dict], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(list(rows), ensure_ascii=False, indent=2), encoding="utf-8")
    return path


EXPORTERS = {"csv": to_csv, "xlsx": to_excel, "json": to_json}


def export(rows: list[dict], path: str | Path) -> Path:
    ext = Path(path).suffix.lower().lstrip(".")
    if ext not in EXPORTERS:
        raise ValueError(f"unsupported output format: .{ext} (use .csv, .xlsx or .json)")
    return EXPORTERS[ext](rows, path)
